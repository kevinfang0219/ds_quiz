from pptx import Presentation
from pptx.util import Inches, Pt

# 建立一個新的簡報物件
prs = Presentation()

# --- Slide 1: 封面 ---
slide = prs.slides.add_slide(prs.slide_layouts[0]) # 標題投影片版面
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "資料結構期末專題\n雜湊表 (Hash Tables) 測驗系統"
subtitle.text = "結合間隔學習法與難易度分級之實作\n報告人：Kevin"

# --- Slide 2: 開發動機與核心亮點 ---
slide = prs.slides.add_slide(prs.slide_layouts[1]) # 標題與內容版面
title, body = slide.shapes.title, slide.placeholders[1]
title.text = "開發動機與核心亮點"
tf = body.text_frame
tf.text = "為什麼要開發這個系統？"
tf.add_paragraph().text = "鎖定核心單元：針對最常考的「雜湊表」進行特訓。"
tf.add_paragraph().text = "解決死背痛點：導入「隨機抽題」與「選項打亂」機制。"
p = tf.add_paragraph()
p.text = "🔥 進階亮點功能："
p.level = 0
tf.add_paragraph().text = "客製化測驗：5 題（快速複習）與 10 題（完整測驗）。"
tf.add_paragraph().text = "間隔學習法：系統記憶錯題，下次測驗優先抽出。"
tf.add_paragraph().text = "難易度視覺化：後台分級，前端即時顯示易/中/難標籤。"

# --- Slide 3: 系統架構與使用技術 ---
slide = prs.slides.add_slide(prs.slide_layouts[1])
title, body = slide.shapes.title, slide.placeholders[1]
title.text = "System Architecture"
tf = body.text_frame
tf.text = "系統架構與使用技術："
tf.add_paragraph().text = "後端框架：Python / Django 6.0.5"
tf.add_paragraph().text = "前端介面：HTML5 / CSS3 (現代科技風 UI、響應式玻璃卡片設計)"
tf.add_paragraph().text = "資料庫系統：SQLite3"
tf.add_paragraph().text = "版本控制：Git / GitHub"

# --- Slide 4: 資料庫關聯設計 ---
slide = prs.slides.add_slide(prs.slide_layouts[1])
title, body = slide.shapes.title, slide.placeholders[1]
title.text = "資料庫設計說明 (ER Model)"
tf = body.text_frame
tf.text = "本系統共設計三個核心資料表："
tf.add_paragraph().text = "1. Question (題目表)：包含題目內容與「難易度」欄位。"
tf.add_paragraph().text = "2. Choice (選項表)：利用 Foreign Key 關聯題目，紀錄正確答案。"
tf.add_paragraph().text = "3. QuizRecord (測驗紀錄表)：記錄秒數，透過 Many-to-Many 關聯紀錄「答錯的題目」，為錯題本功能核心。"

# --- Slide 5: 系統 Live Demo ---
slide = prs.slides.add_slide(prs.slide_layouts[1])
title, body = slide.shapes.title, slide.placeholders[1]
title.text = "系統 Live Demo 🚀"
tf = body.text_frame
tf.text = "展示重點路線："
tf.add_paragraph().text = "1. 展示首頁與動態難易度徽章"
tf.add_paragraph().text = "2. 展示隨機抽題與選項洗牌機制"
tf.add_paragraph().text = "3. 答題回饋 (綠色正確/紅色錯誤)"
tf.add_paragraph().text = "4. 🔥 殺手鐧展示：結算後再次測驗，展示「錯題優先抽出」機制"

# --- Slide 6: 結論與未來展望 ---
slide = prs.slides.add_slide(prs.slide_layouts[1])
title, body = slide.shapes.title, slide.placeholders[1]
title.text = "Conclusion & Future Work"
tf = body.text_frame
tf.text = "系統結論："
tf.add_paragraph().text = "成功實作基本要求，並超前部署多項進階功能。"
p = tf.add_paragraph()
p.text = "未來可擴充方向："
p.level = 0
tf.add_paragraph().text = "導入會員系統，實作「班級排行榜」。"
tf.add_paragraph().text = "加入視覺化的 JS 倒數計時器。"
tf.add_paragraph().text = "擴充更多單元 (如 Tree, Graph)。"

# 儲存簡報
prs.save('期末專題簡報_HashTables.pptx')
print("🎉 成功生成簡報！請查看資料夾中的 '期末專題簡報_HashTables.pptx'")